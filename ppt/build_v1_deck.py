# -*- coding: utf-8 -*-
"""
依据 V1_中文研究逻辑与PPT母稿.md 第 8 节的 12 页提纲生成汇报 PPT。
- 视觉基底沿用 moban.pptx 的母版/主题（#025483 / 微软雅黑 / 校训页脚）
- 配图取自 Conference-LaTeX-template_10-17-19/figs/
- 严格执行母稿的表达约定：标注 MC 次数与 83.7% 口径，区分「模型内证明 / 实验观测 / 机制解释」
输出：D:\\Desktop\\V1_目标特定融合与通信受限证据选择.pptx
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = r"D:\Desktop\moban.pptx"
FIG_DIR = r"D:\Desktop\conference\Conference-LaTeX-template_10-17-19\figs"
OUT = r"D:\Desktop\V1_目标特定融合与通信受限证据选择.pptx"

# ---------------- 设计系统 ----------------
BLUE = RGBColor(0x02, 0x54, 0x83)
DEEP = RGBColor(0x00, 0x20, 0x60)
GREY = RGBColor(0x44, 0x54, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xEE, 0xF3, 0xF8)
EDGE = RGBColor(0xC9, 0xD8, 0xE6)
ORANGE = RGBColor(0xC0, 0x6A, 0x2E)   # 实验观测标记
RED = RGBColor(0x99, 0x3A, 0x3A)      # 不支持 / 边界标记
CN, EN, MATH = "微软雅黑", "Arial", "Cambria Math"

SW, SH = 13.333, 7.5
ML, MR = 0.6, 0.6
CW = SW - ML - MR
TOP, BOT = 1.55, 6.55


# ---------------- 基础工具 ----------------
def set_ea(run, name=CN):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def style(run, size=None, bold=None, color=None, cn=CN, en=EN):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    f.name = en
    set_ea(run, cn)
    return run


def set_text(tf, lines, size=14, bold=None, color=GREY, align=PP_ALIGN.LEFT,
             cn=CN, en=EN, space_after=6, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass
    tf.clear()
    if isinstance(lines, str):
        lines = [lines]
    for i, item in enumerate(lines):
        if isinstance(item, str):
            item = (item, size, bold, color)
        txt, sz, bd, cl = item
        p = tf.paragraphs[0] if (i == 0 and len(tf.paragraphs)) else tf.add_paragraph()
        p.alignment = align
        if space_after is not None:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = txt
        style(r, sz, bd, cl, cn, en)
    return tf


def put_text(slide, x, y, w, h, lines, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    set_text(box.text_frame, lines, **kw)
    return box


def rect(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line_w=1.0, adj=0.05):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    sp.text_frame.word_wrap = True
    return sp


def find_shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def replace_text(slide, name, lines, **kw):
    sh = find_shape(slide, name)
    if sh is None:
        raise KeyError(f"未找到形状：{name}")
    set_text(sh.text_frame, lines, **kw)
    return sh


def new_slide(prs, layout_name="1_自定义版式"):
    for l in prs.slide_masters[0].slide_layouts:
        if l.name == layout_name:
            return prs.slides.add_slide(l)
    return prs.slides.add_slide(prs.slide_masters[0].slide_layouts[2])


def set_title(slide, text, sub=None):
    sh = None
    for p in slide.placeholders:
        if p.placeholder_format.idx == 10:
            sh = p
            break
    if sh is None:
        sh = put_text(slide, 1.68, 0.62, 9.4, 0.5, "")
    sh.left, sh.top = Inches(1.68), Inches(0.60)
    sh.width, sh.height = Inches(9.6), Inches(0.52)
    set_text(sh.text_frame, text, size=23, bold=True, color=BLUE, space_after=0)
    if sub:
        put_text(slide, 1.72, 1.16, 10.2, 0.3, sub, size=11.5, color=GREY)
    return sh


def note(slide, text):
    put_text(slide, 2.60, 6.74, 9.4, 0.3, text, size=9, color=GREY)


def badge(slide, x, y, kind):
    """三种证据标记：模型内证明 / 实验观测 / 机制解释。"""
    cfg = {
        "证明": (BLUE, WHITE),
        "实验": (ORANGE, WHITE),
        "解释": (GREY, WHITE),
        "边界": (CARD, RED),
    }[kind]
    sp = rect(slide, x, y, 0.58, 0.30, fill=cfg[0], shape=MSO_SHAPE.ROUNDED_RECTANGLE,
              adj=0.3)
    set_text(sp.text_frame, kind, size=10, bold=True, color=cfg[1],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    return sp


def card(slide, x, y, w, h, head, body, head_h=0.46, body_size=13,
         tag=None, head_size=15):
    rect(slide, x, y, w, h, fill=CARD, line=EDGE)
    hd = rect(slide, x, y, w, head_h, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    set_text(hd.text_frame, head, size=head_size, bold=True, color=WHITE,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    if tag:
        badge(slide, x + w - 0.68, y + 0.08, tag)
    tb = put_text(slide, x + 0.22, y + head_h + 0.14, w - 0.44, h - head_h - 0.28,
                  body, size=body_size, color=GREY, space_after=7)


def pic(slide, name, x, w, y):
    p = os.path.join(FIG_DIR, name)
    sh = slide.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))
    sh.shadow.inherit = False
    return sh


# ================= 打开母版文件，保留封面与尾页 =================
prs = Presentation(SRC)
sldIdLst = prs.part._element.sldIdLst
ids = list(sldIdLst)
for e in ids[1:8]:                      # 删掉中间 7 页骨架，只留封面与尾页
    prs.part.drop_rel(e.get(qn("r:id")))
    sldIdLst.remove(e)
cover, ending = list(prs.slides)[0], list(prs.slides)[1]

# ---------- 封面 ----------
replace_text(cover, "矩形 27", "目标特定融合与通信受限证据选择", size=40, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for para in find_shape(cover, "矩形 27").text_frame.paragraphs:
    for r in para.runs:
        r.font.color.rgb = WHITE
put_text(cover, 2.0, 4.38, 9.33, 0.5,
         "多无人机 OTFS-ISAC 协同感知中的证据选择", size=17, color=WHITE,
         align=PP_ALIGN.CENTER)
rep = replace_text(cover, "文本框 21", "汇报人：XXX", size=14, color=BLUE,
                   align=PP_ALIGN.CENTER)
rep.left, rep.width = Inches(4.6), Inches(4.13)

# ---------- 尾页 ----------
replace_text(ending, "矩形 6", "THANK YOU", size=41.25, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for para in find_shape(ending, "矩形 6").text_frame.paragraphs:
    for r in para.runs:
        r.font.color.rgb = WHITE


# ================= 第 1 页：研究对象 / 科学问题 / 一句结果 =================
s = new_slide(prs)
set_title(s, "融合位置如何影响协同感知的报告效率")
cw, gap = 3.83, 0.32
card(s, ML + 0 * (cw + gap), 1.85, cw, 4.1, "研究对象",
     [("多无人机协同感知把回波观测分布到多个接收节点，检测证据需要在某个节点汇合。", 13, False, GREY),
      ("任务是已有上游预测结果的目标确认或重检测。", 13, False, GREY)])
card(s, ML + 1 * (cw + gap), 1.85, cw, 4.1, "科学问题",
     [("能否通过目标相关的融合位置配置，以及显式考虑该位置的证据选择，",
       13, False, GREY),
      ("提高报告通信效率？", 13, True, BLUE)])
card(s, ML + 2 * (cw + gap), 1.85, cw, 4.1, "一句结果",
     [("1000 次配对试验：检测概率差异 −0.0021（未显著），", 13, False, GREY),
      ("远程报告 4.595 → 0.751 条，", 15, True, BLUE),
      ("固定包长串行模型下负载与时延降低 83.7%。", 13, False, GREY)],
     tag="实验")
note(s, "注：83.7% 相对 Sensing-SINR 基线；同融合图、匹配每目标观测数、未匹配通信预算。")


# ================= 第 2 页：观测位置 ≠ 融合位置 =================
s = new_slide(prs)
set_title(s, "观测产生位置与融合位置可以不同",
          "软统计量在接收节点 j 形成，却可能要在另一个节点 f_q 汇合")
left = rect(s, ML, 1.85, 6.55, 3.75, fill=CARD, line=EDGE)
put_text(s, ML + 0.25, 1.98, 6.05, 0.35, "信息路径", size=14, bold=True, color=BLUE)
nodes = [("i", "发射\n无人机"), ("q", "目标"), ("j", "接收\n无人机"), ("f_q", "融合\n节点")]
nx = ML + 0.30
for k, (lab, cap) in enumerate(nodes):
    w = 1.30 if k < 3 else 1.35
    nd = rect(s, nx, 2.60, w, 0.80, fill=WHITE, line=BLUE, line_w=1.5, adj=0.15)
    set_text(nd.text_frame, lab, size=17, bold=True, color=BLUE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             en=MATH, cn=CN, space_after=0)
    put_text(s, nx - 0.15, 3.45, w + 0.30, 0.55, cap.replace("\n", ""),
             size=11, color=GREY, align=PP_ALIGN.CENTER)
    if k < 3:
        rect(s, nx + w + 0.03, 2.85, 0.36, 0.30, fill=BLUE,
             shape=MSO_SHAPE.RIGHT_ARROW)
    nx += w + 0.42
band1 = rect(s, ML + 0.30, 4.15, 2.55, 0.45, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
set_text(band1.text_frame, "i→q→j 决定感知观测", size=12, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
band2 = rect(s, ML + 3.40, 4.15, 2.85, 0.45, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
set_text(band2.text_frame, "j→f_q 决定报告条件", size=12, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
put_text(s, ML + 0.30, 4.75, 5.95, 0.75,
         "改变融合节点不会在当前固定观测模型中增强同一回波，",
         size=12, color=GREY)
put_text(s, ML + 0.30, 5.02, 5.95, 0.75,
         "却会改变报告路径及候选观测能否参与融合。", size=12, bold=True, color=BLUE)

c1 = rect(s, 7.53, 1.85, 5.40, 1.78, fill=CARD, line=EDGE)
h1 = rect(s, 7.53, 1.85, 5.40, 0.46, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
set_text(h1.text_frame, "本地证据　j = f_q", size=14, bold=True, color=WHITE,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space_after=0, en=MATH)
put_text(s, 7.75, 2.42, 4.96, 1.05,
         [("χ = 1，无需发送报告包，报告成本为零。", 13, False, GREY),
          ("不意味着感知与处理没有成本。", 12, False, GREY)], space_after=6)
c2 = rect(s, 7.53, 3.82, 5.40, 1.78, fill=CARD, line=EDGE)
h2 = rect(s, 7.53, 3.82, 5.40, 0.46, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
set_text(h2.text_frame, "远程报告　j ≠ f_q", size=14, bold=True, color=WHITE,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space_after=0, en=MATH)
put_text(s, 7.75, 4.39, 4.96, 1.05,
         [("统计量经 j→f_q 交付，必须满足通信距离、", 13, False, GREY),
          ("速率与成功率等可行性要求，并占用报告成本。", 13, False, GREY)],
         space_after=6)
note(s, "注：f_q 为目标 q 的融合节点，不同目标可以共享同一架无人机。")


# ================= 第 3 页：选择要同时考虑感知与报告 =================
s = new_slide(prs)
set_title(s, "选择证据需要同时考虑感知与报告",
          "强回波可能来自远程节点，弱观测可能已在融合节点本地可用")
for i, (t, b) in enumerate([
    ("感知质量", "回波强弱决定单条观测本身能提供的检测信息。"),
    ("交付可靠度", "远程报告经有限码长链路，成功率为 χ，不一定送达。"),
    ("报告代价", "每条远程报告占用固定包长与信道使用，本地为零。"),
]):
    card(s, ML + i * (cw + gap), 1.85, cw, 1.95, t, b, head_size=14, body_size=13)
ex = rect(s, ML, 4.05, CW, 1.75, fill=CARD, line=EDGE)
eh = rect(s, ML, 4.05, CW, 0.44, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
set_text(eh.text_frame, "算例（说明报告数公式，不是新增实验）", size=13, bold=True,
         color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
put_text(s, ML + 0.28, 4.62, CW - 0.56, 1.05,
         [("某目标选中 4 条观测：3 条产生于无人机 A、1 条产生于 B，报告均可行。", 13, False, GREY),
          ("融合在 A → 1 次远程报告 = 640 bit；融合在未产生这些观测的 C → 4 次 = 2560 bit。",
           14, True, BLUE)], space_after=7)
note(s, "注：这是可能出现的价值排序变化，不能表述为「较弱的本地观测总是优于较强的远程观测」。")


# ================= 第 4 页：预测目标位置作为配置依据 =================
s = new_slide(prs)
set_title(s, "预测目标位置提供低复杂度配置依据",
          "先配置融合节点，再在该位置决定的报告图上选择观测")
card(s, ML, 1.85, 5.90, 2.55, "为什么可以使用预测位置",
     [("任务是已有上游预测结果的目标确认或重检测；", 13, False, GREY),
      ("规划使用预测均值与协方差，不使用当前目标真值或已实现回波。", 13, False, GREY)],
     head_size=14, body_size=13)
fbox = rect(s, 6.83, 1.85, 5.90, 2.55, fill=CARD, line=EDGE)
fh = rect(s, 6.83, 1.85, 5.90, 0.46, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
set_text(fh.text_frame, "最近预测目标规则（nearest-target）", size=14, bold=True,
         color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
put_text(s, 6.83, 2.62, 5.90, 0.75,
         "f_q = arg min_m ‖ p̂_q − p_m^UAV ‖₂", size=21, bold=True, color=DEEP,
         align=PP_ALIGN.CENTER, en=MATH)
put_text(s, 7.05, 3.42, 5.46, 0.85,
         "候选集合随融合位置变化：E_q(f_q) = {(i,j,q)：DD 有效，且本地或报告链路可行}。",
         size=12.5, color=GREY)
warn = rect(s, ML, 4.62, CW, 1.68, fill=CARD, line=EDGE)
badge(s, ML + 0.24, 4.80, "边界")
put_text(s, ML + 0.98, 4.78, CW - 1.20, 1.35,
         [("nearest-target 是证据空间分布的代理，不是最优位置定理：距离衰减之外，",
           13, False, GREY),
          ("发射几何、干扰与 DD 泄漏同样起作用；预测状态与无人机状态的分发成本未完整建模，",
           13, False, GREY),
          ("因此不能称为「零额外控制开销」。", 13, True, RED)], space_after=6)


# ================= 第 5 页：三模块紧凑模型 =================
s = new_slide(prs)
set_title(s, "紧凑模型连接几何、统计量与报告",
          "每个模型项都服务于「选择哪几条观测」，不做多余延伸")
mods = [
    ("双基地 DD 几何",
     [("时延取两段距离和除以光速，多普勒含目标与两端相对运动；", 12.5, False, GREY),
      ("粗级 sinc 近似、精细级有限周期网格局部能量。", 12.5, False, GREY),
      ("不宣称波形创新或最佳窗口。", 12, True, RED)]),
    ("有限码长报告成功率",
     [("正态近似把通信 SINR 映射为成功概率 χ；", 12.5, False, GREY),
      ("失败报告采用高斯替代，方差系数 a = 9。", 12.5, False, GREY),
      ("不是实测误包曲线或精确丢包似然。", 12, True, RED)]),
    ("固定报告成本",
     [("一条远程报告固定占用 640 bit 与 2048 个信道使用；", 12.5, False, GREY),
      ("不聚合、不重传。", 12.5, False, GREY),
      ("成本只随报告条数变化。", 12, True, RED)]),
]
for i, (t, b) in enumerate(mods):
    card(s, ML + i * (cw + gap), 1.85, cw, 2.85, t, b, head_size=14, body_size=12.5)
fml = rect(s, ML, 5.00, CW, 1.05, fill=WHITE, line=EDGE)
set_text(fml.text_frame,
         "B_report = 640 · N_rem　　T_report = N_rem · n / B_c",
         size=17, bold=True, color=DEEP, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, en=MATH, space_after=0)
note(s, "注：链路条件改变可靠度与可行性，报告条数改变成本；不能用「链路缩短使每个包更快」解释当前时延结果。")


# ================= 第 6 页：目标函数 =================
s = new_slide(prs)
set_title(s, "检测效用与报告价格共同决定选择",
          "规划在 belief 侧预测交付后的检测概率，并对远程报告收费")
fbox = rect(s, ML, 1.90, 5.55, 2.30, fill=WHITE, line=BLUE, line_w=1.5)
set_text(fbox.text_frame,
         "F(S) = U( P̂_D(S) ) − λ_c · Σ_{e∈S} c_e",
         size=20, bold=True, color=DEEP, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, en=MATH, space_after=0)
put_text(s, ML, 4.28, 5.55, 0.9,
         "名义价格 λ_c = 0.005 / ms；c_e 以毫秒计，本地证据代价为零。",
         size=12.5, color=GREY, align=PP_ALIGN.CENTER)
pts = [
    ("效用 U 的三项机制",
     "在设计要求处截断；用 soft-min 照顾弱目标；用二次缺口项惩罚未满足要求的目标。"),
    ("直觉",
     "达到设计要求后，继续增加该目标观测的收益递减；选择器把有限容量转向尚弱的目标。"),
    ("边界",
     "两项弱目标机制的独立必要性尚无完整消融，不能各自包装成创新；也无法把 83.7% 拆给各效用项。"),
]
y = 1.90
for t, b in pts:
    card(s, 6.45, y, 6.28, 1.42, t, b, head_size=13.5, body_size=12)
    y += 1.58
note(s, "注：目标函数使用检测器一致的矩近似预测检测概率，不是精确最优检测。")


# ================= 第 7 页：C2F =================
s = new_slide(prs)
set_title(s, "C2F 将精细计算集中到候选前沿",
          "粗级筛出短名单，只对进入前沿的候选做精细评估")
steps = ["① 粗级 DD\n估计与效用", "② 动态构建\n候选前沿", "③ 前沿内\n精细 DD",
         "④ 重置并在精细\n估计上重选", "⑤ 贪心加入\n正增益候选"]
bw, gp = 2.10, 0.35
for i, t in enumerate(steps):
    x = ML + 0.10 + i * (bw + gp)
    b = rect(s, x, 1.90, bw, 1.25, fill=CARD, line=EDGE)
    set_text(b.text_frame, t, size=12.5, bold=(i in (1, 3)), color=BLUE if i in (1, 3) else GREY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(steps) - 1:
        rect(s, x + bw + 0.02, 2.36, gp - 0.04, 0.28, fill=BLUE,
             shape=MSO_SHAPE.RIGHT_ARROW)
put_text(s, ML + 0.10, 3.28, 11.9, 0.3,
         "单目标前沿上限 J = 25；每步加入满足约束且具有最大正边际增益的候选，无正增益或达到上限即停止。",
         size=12, color=GREY)
fml = rect(s, ML, 3.70, 4.20, 0.85, fill=WHITE, line=EDGE)
set_text(fml.text_frame, "DD 计算工作量 ≈ E·C_c + r·C_f，　r ≤ E",
         size=16, bold=True, color=DEEP, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, en=MATH, space_after=0)
kpis = [("61.3", "平均精细评估次数"), ("860.2", "可行候选数"), ("92.9%", "工作量计数减少")]
for i, (v, cap) in enumerate(kpis):
    x = ML + 4.60 + i * 2.60
    rect(s, x, 3.70, 2.35, 0.85, fill=CARD, line=EDGE)
    put_text(s, x, 3.80, 2.35, 0.45, v, size=20, bold=True, color=BLUE,
             align=PP_ALIGN.CENTER)
    put_text(s, x, 4.22, 2.35, 0.3, cap, size=10.5, color=GREY, align=PP_ALIGN.CENTER)
warn = rect(s, ML, 4.80, CW, 1.42, fill=CARD, line=EDGE)
badge(s, ML + 0.24, 4.98, "边界")
put_text(s, ML + 0.98, 4.96, CW - 1.20, 1.10,
         [("工作量式不含全部贪心评分、缓存与共享表构建开销；当前没有遗漏候选效用界，",
           12.5, False, GREY),
          ("因此 C2F 不保证无损筛选，也没有全局近似比保证。", 12.5, True, RED)],
         space_after=6)


# ================= 第 8 页：主比较 =================
s = new_slide(prs)
set_title(s, "同一融合图下，远程报告减少 83.7%",
          "MC = 1000 次配对试验，共同场景与物理模型；本页只检验选择策略")
main_kpi = [("−0.0021", "检测概率配对差\n95%CI [−0.0051, 0.0009]"),
            ("12.101", "平均选中观测数\n两种方法完全相同"),
            ("0.751 ← 4.595", "平均远程报告数\nV1 ← Sensing-SINR"),
            ("−83.7%", "负载与时延\n0.481 kbit / 0.801 ms")]
kw, kgap = 2.72, 0.4
for i, (v, cap) in enumerate(main_kpi):
    x = ML + i * (kw + kgap)
    rect(s, x, 1.62, kw, 1.45, fill=CARD, line=EDGE)
    rect(s, x, 1.62, kw, 0.10, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    put_text(s, x + 0.06, 1.80, kw - 0.12, 0.55, v, size=19, bold=True, color=BLUE,
             align=PP_ALIGN.CENTER)
    put_text(s, x + 0.06, 2.38, kw - 0.12, 0.6, cap, size=10, color=GREY,
             align=PP_ALIGN.CENTER)
pic(s, "fig2_v1_main_comparison.png", 1.17, 11.0, 3.30)
note(s, "注：Sensing-SINR 与 V1 使用相同融合图、可行性检查与匹配的每目标观测数；两者未匹配远程通信预算。")


# ================= 第 9 页：融合消融 =================
s = new_slide(prs)
set_title(s, "改变融合规则会改变最终报告需求",
          "MC = 100 次／每种规则，属机制诊断，不可与 1000 次主比较混算")
rows, cols = 4, 5
tbl = s.shapes.add_table(rows, cols, Inches(ML), Inches(1.62), Inches(CW), Inches(1.95)).table
data = [["融合规则", "平均检测概率", "平均选中观测数", "平均远程报告数", "串行报告时延"],
        ["最大入速率", "0.919", "19.89", "6.65", "7.093 ms"],
        ["最大最小速率", "0.969", "15.61", "6.40", "6.827 ms"],
        ["最近预测目标", "0.982", "12.24", "0.74", "0.789 ms"]]
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        head = (r == 0)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
        run = p.add_run()
        run.text = data[r][c]
        style(run, size=12.5, bold=head, color=WHITE if head else GREY)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if head else (WHITE if r % 2 else CARD)
pic(s, "fig3_v1_fusion_ablation.png", 2.17, 9.0, 3.80)
note(s, "注：更换位置会同时改变可行集合与最终选中集合，不能解释为固定一组观测后的纯交付效应；时延仍为串行 MAC。")


# ================= 第 10 页：理论解释 =================
s = new_slide(prs)
set_title(s, "理论与消融解释本地证据的价值",
          "三层结论各自成立，但都不能扩大为保证")
theorems = [
    ("① 固定观测的本地交付优势", "证明",
     "d(χ) = χ²δ² / ( v₀ [ a + (1−a)χ ] )",
     "在规定条件与正分母下随 χ 非减：同一观测本地交付有利于该代理信息指标，"
     "不说明 nearest-target 对所有观测组合最优。"),
    ("② 最优检测能力的条件性支配", "证明",
     "Q_{h,χ} = χ·P_h + (1−χ)·G",
     "较低 χ 的输出可由较高 χ 的输出通过随机保留或重新替代构造，因此较高可靠度的最优 ROC 包络不会更差；"
     "不证明当前线性检测器或矩近似具有同样单调性。"),
    ("③ 固定观测集合的报告数界", "证明",
     "N_rem,q(m) = |S_q| − a_qm ≥ |S_q| − max_m a_qm",
     "满足可行性且达到右侧界的融合节点，对这个固定集合报告最少；"
     "改变位置后重新选择观测已是另一个问题。"),
]
for i, (t, tag, fml, body) in enumerate(theorems):
    x = ML + i * (cw + gap)
    card(s, x, 1.80, cw, 4.30, t, "", head_h=0.46, tag=tag)
    fb = rect(s, x + 0.18, 2.42, cw - 0.36, 0.72, fill=WHITE, line=EDGE)
    set_text(fb.text_frame, fml, size=12.5, bold=True, color=DEEP,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, en=MATH, space_after=0)
    put_text(s, x + 0.22, 3.28, cw - 0.44, 2.65, body, size=11.5, color=GREY,
             space_after=4)
note(s, "注：图中「证明」＝模型内成立；实验支持来自主比较与融合消融；几何解释尚未分离证明全部收益来自距离因素。")


# ================= 第 11 页：计算收益与边界 =================
s = new_slide(prs)
set_title(s, "计算收益明确，适用边界仍需交代")
card(s, ML, 1.85, 5.90, 3.95, "计算收益",
     [("· 主实验：精细评估 61.3 次 vs 860.2 个可行候选，减少 92.9%。", 12.5, False, GREY),
      ("· 独立 200 次 full-refinement 控制：61.06 vs 853.65 次，减少 92.85%；"
       "检测概率 0.980 vs 0.973（后者不是 oracle）。", 12.5, False, GREY),
      ("· 独立 30 场景计时：中位数 0.556 s vs 2.032 s，约 3.66 倍；"
       "排除共享表与预计算，不是端到端加速。", 12.5, False, GREY),
      ("· 矩预测与采样检测概率最大差约 0.0447；精确单报告混合分布最大差约 0.00184。",
       12.5, False, GREY)],
     head_size=14, body_size=12.5, tag="实验")
card(s, 6.83, 1.85, 5.90, 3.95, "适用边界",
     [("· 位置误差 50 m → 500 m：检测概率约 0.99 → 0.900。", 12.5, False, GREY),
      ("· 部署区域边长 2500 m → 5500 m：约 0.992 → 0.932，"
       "远程报告 0.05 → 2.67。", 12.5, False, GREY),
      ("· 核验工具尚未替换默认 C2F 评分，不代表实际检测率被提高。", 12.5, False, GREY),
      ("· 上述扫描限定方法适用边界，不构成统一缩放定律。", 12.5, True, RED)],
     head_size=14, body_size=12.5, tag="边界")
note(s, "注：C2F 页使用主稿 92.9% 的精细评估计数口径；3.66 倍来自另一组 30 场景限定计时。")


# ================= 第 12 页：总结 =================
s = new_slide(prs)
set_title(s, "主线是位置感知的证据选择")
card(s, ML + 0 * (cw + gap), 1.85, cw, 4.05, "已证实的结果",
     [("· 融合位置改变报告属性、可行性与成本（模型内成立）。", 12.5, False, GREY),
      ("· 本地交付有条件的检测信息优势（条件性理论结论）。", 12.5, False, GREY),
      ("· 给定融合图后 V1 显著减少远程报告：0.751 vs 4.595，"
       "同观测数下检测差异未显著。", 12.5, False, GREY)],
     head_size=14, body_size=12.5, tag="实验")
card(s, ML + 1 * (cw + gap), 1.85, cw, 4.05, "尚未证明的部分",
     [("· nearest-target 接近联合最优：尚无精确联合基准。", 12.5, False, GREY),
      ("· C2F 安全保留最优候选：尚无遗漏候选界。", 12.5, False, GREY),
      ("· 检测统计等效或普遍非劣：未预设等效界、未完成对应检验。", 12.5, False, GREY),
      ("· 不能把 83.7% 外推为全系统时延降低。", 12.5, True, RED)],
     head_size=14, body_size=12.5, tag="边界")
card(s, ML + 2 * (cw + gap), 1.85, cw, 4.05, "下一步",
     [("· 小规模精确「位置—子集」联合基准。", 12.5, False, GREY),
      ("· 同通信预算下的性能比较。", 12.5, False, GREY),
      ("· 以上为计划，不列入已完成贡献。", 12.5, True, RED)],
     head_size=14, body_size=12.5)
note(s, "收束：目标特定融合目的节点 + 该图上的检测器一致选择，构成本文主线。")


# ================= 备份页：名义参数与未计入开销 =================
s = new_slide(prs)
set_title(s, "备份：名义参数与未计入开销", "主讲不展开，用于回答追问")
params = [["无人机 / 目标", "15 架 / 10 个"], ["部署区域", "4000 m × 4000 m"],
          ["预测标准差", "位置 150 m，速度 15 m/s"],
          ["每目标观测上限", "6 条（总计最多 60 条）"],
          ["设计检测概率 / 虚警", "0.95 / 0.05"],
          ["功率分配 ρ", "0.8（感知 ρP，报告 (1−ρ)P）"],
          ["失败报告方差系数 a", "9"], ["报告价格 λ_c", "0.005 / ms"],
          ["前沿上限 J", "25"], ["单条远程报告", "640 bit，2048 信道使用"]]
rows = len(params) + 1
tbl = s.shapes.add_table(rows, 2, Inches(ML), Inches(1.78), Inches(5.90), Inches(4.55)).table
tbl.cell(0, 0).text = ""
head = ["参数", "取值"]
for c, h in enumerate(head):
    cell = tbl.cell(0, c)
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = h
    style(run, size=12, bold=True, color=WHITE)
    cell.fill.solid()
    cell.fill.fore_color.rgb = BLUE
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
for r, (k, v) in enumerate(params, start=1):
    for c, txt in enumerate((k, v)):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = txt
        style(run, size=11.5, bold=False, color=GREY)
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if r % 2 else CARD
card(s, 6.83, 1.78, 5.90, 4.55, "未计入的开销与常见追问",
     [("· 控制分发、包头、重传、聚合、处理与端到端协调开销均未计入。", 12, False, GREY),
      ("· 为什么不是简单选最近节点？最近规则只是位置配置的一个实例；"
       "系统问题是融合目的节点会改变后续报告条件。", 12, False, GREY),
      ("· Sensing-SINR 检测略高？它优先感知质量且不计入报告价格；"
       "当前差异未显著，本文强调该条件下的通信效率。", 12, False, GREY),
      ("· 本地「免费」是否人为制造优势？零仅指该报告包的传输成本。", 12, False, GREY)],
     head_size=14, body_size=12)


# ================= 收尾：清空硬编码页码 + 重排 =================
for sl in prs.slides:
    for ph in sl.placeholders:
        if ph.placeholder_format.idx == 4:
            set_text(ph.text_frame, "")

order = [0] + list(range(2, 15)) + [1]
ids = list(prs.part._element.sldIdLst)
assert len(ids) == 15, f"共 {len(ids)} 页，与预期 15 不符"
for e in ids:
    prs.part._element.sldIdLst.remove(e)
for i in order:
    prs.part._element.sldIdLst.append(ids[i])

prs.save(OUT)
print("saved:", OUT, "| 页数:", len(ids))
