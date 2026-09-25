"""Build the living Chinese thesis research report for direction 4."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "多无人机OTFS通感一体化共享状态失配与校正阶段报告_总体框架修订版.docx"
RETENTION = ROOT / "studies" / "direction4" / "data" / "shared_mismatch_retention" / "summary.json"
GATE1 = ROOT / "studies" / "direction4" / "data" / "gate1" / "summary.json"
STAGEB_SMOKE = (ROOT / "studies" / "direction4" / "data" /
                "stageb_one_scene_r1" / "summary.json")


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, first_line=True, before=0, after=0, line=1.5):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74) if first_line else None
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = line
    return paragraph


def add_body(doc, text, first_line=True):
    p = style_paragraph(doc.add_paragraph(), first_line=first_line)
    set_run_font(p.add_run(text))
    return p


def add_formula(doc, latex, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    text = latex if number is None else f"{latex}    ({number})"
    set_run_font(p.add_run(text), italic=True)
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 12}.get(level, 12), bold=True)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.37)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(item))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, "D9E2F3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(str(header)), bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(str(value)))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_status_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    shade(table.cell(0, 0), "EEF3F8")
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(text), bold=True, color=(31, 78, 121))
    return table


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def summary_value(retention, radius, count=6):
    row = next(x for x in retention["summary"]
               if x["error_radius_m"] == radius and x["receiver_count"] == count)
    return row["retention_mean"], row["retention_ci95_low"], row["retention_ci95_high"]


def build():
    retention = json.loads(RETENTION.read_text(encoding="utf-8"))
    gate1 = json.loads(GATE1.read_text(encoding="utf-8"))
    stageb = (json.loads(STAGEB_SMOKE.read_text(encoding="utf-8"))
              if STAGEB_SMOKE.exists() else None)
    eval_by_name = {x["method"]: x for x in gate1["evaluation"]}

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    inherited_border = title_ppr.find(qn("w:pBdr"))
    if inherited_border is not None:
        title_ppr.remove(inherited_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(24)
    set_run_font(title.add_run("多无人机 OTFS 通感一体化系统中的共享目标状态失配与交叉拟合校正"), 20, True)
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("硕士学位论文阶段研究报告"), 16, True)
    doc.add_paragraph().paragraph_format.space_after = Pt(96)
    meta = [
        ("研究对象", "6 架 UAV 协同感知 3 个运动目标"),
        ("核心问题", "共享目标状态偏差引起的多节点检测失配及其数据隔离校正"),
        ("报告版本", "阶段稿 2026 年 9 月"),
    ]
    add_table(doc, ["项目", "内容"], meta, [3.5, 10.5])
    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_body(doc, "本报告围绕多无人机正交时频空间通感一体化系统中的共享目标状态失配展开。现有系统在 6 架无人机协同探测 3 个运动目标时曾出现最差检测概率约为 0.2 的异常结果。审计表明，该现象的主要原因不是节点数量不足、通信报告预算或直达径干扰消除失效，而是所有接收机共同使用了严重偏移的目标状态字典。公共模型偏差会同时削弱各节点的匹配滤波证据，因此增加节点只能降低随机波动，不能自动恢复被错误字典丢失的信息。")
    add_body(doc, "本文将研究收缩为两个递进问题。问题 1 建立共享状态失配下的多节点 OTFS 感知模型，并用字典子空间信息保留率解释检测性能退化。问题 2 在相邻短时相干处理区间内，将状态估计数据与检测数据分离，构建共享偏移的交叉拟合最大后验校正方法。首轮 12 场景机制实验显示，在位置先验标准差为 150 m 时，六节点信息保留率随归一化误差由 0 降至 2.67，从 1.000 降至 0.039；节点数从 1 增至 6 并未提高平均保留率。既有 150 m 探索性检测实验中，共享交叉拟合相对名义字典的 AUC 提高 0.0725、检测概率提高 0.15，但置信区间和旧几何约束决定了该结果仍需在 20 m 最短节点间距下复验。")
    add_body(doc, "本阶段的结论是：研究主线应从队形或干扰消除的泛化优化，收缩为共享状态失配的机制解释与低偏差校正。OTFS 在本文中的意义来自运动目标产生的时延与多普勒联合表征；场景采用短时块静态位置和非零速度，而不是完全静止目标。后续实验将按机制验证、单目标检测曲线、多目标泛化和队形扩展的顺序推进。")
    p = doc.add_paragraph()
    set_run_font(p.add_run("关键词："), bold=True)
    set_run_font(p.add_run("OTFS；通感一体化；多无人机协同感知；共享状态失配；交叉拟合；目标检测"))

    add_heading(doc, "术语与证据口径", 1)
    add_table(doc, ["规范术语", "本文定义", "使用规则"], [
        ("共享目标状态失配", "多个接收机共同使用同一偏移目标状态构造字典", "不写成独立测量噪声"),
        ("名义方法 nominal", "不更新先验状态，直接检测", "正式基线"),
        ("共享交叉拟合 shared cross fit", "一折估状态，另一折检测，再交换", "拟议方法"),
        ("oracle", "使用真值状态构造字典", "仅作性能上界"),
        ("归一化误差 r", "r 等于位置偏差范数除以先验标准差", "跨场景主横轴"),
        ("信息保留率 eta", "真值字典在信念字典子空间内的能量比例", "机制指标，不等同于检测概率"),
        ("几何敏感矩阵 Jjq", "共享状态偏差到节点时延、多普勒和角度偏差的一阶映射", "理论分析量"),
        ("状态信息矩阵 Iq", "先验信息与多节点几何信息之和", "可辨识性诊断量，待实验验证"),
    ], [3.2, 7.2, 4.0])
    add_heading(doc, "1 研究背景及意义", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "通信感知一体化（Integrated Sensing and Communication，ISAC）正在由频谱、硬件和波形层面的功能复用，逐步发展为通信、感知与网络状态相互辅助的联合处理体系。Liu 等将这一演进概括为通信辅助感知、感知辅助通信以及网络化感知，说明位置、速度、信道与环境状态正在从系统输出进一步转化为接收处理和资源控制的输入[1]。在低空网络中，无人机具有部署灵活、视距传播概率较高和观测几何可调等特点，因而既可以作为通信节点，也可以形成分布式感知网络。")
    add_body(doc, "面向无人机和高速运动目标，较大的多普勒扩展会削弱传统时频域处理的稳定性。OTFS 将信号组织到时延—多普勒（delay-Doppler，DD）域，使传播路径能够由具有物理含义的时延、多普勒及角度参数表征。Gaudio 等证明了 OTFS 在联合通信与雷达参数估计中的可行性[2]。此后，相关研究的重点不再只是回答“OTFS 能否感知”，而是转向目标参数与连续物理状态如何被更准确、低复杂度地恢复。本文采用短时块静态位置和非零速度模型：一个相干处理区间内固定几何，在相邻区间保留运动状态。因此，OTFS 的作用来自运动目标的时延—多普勒联合表征，而不是为完全静止场景增加一种装饰性波形。")
    add_body(doc, "近五年的 OTFS 接收算法形成了清晰的演进链。固定 DD 网格上的匹配估计首先发展为利用稀疏结构的 OMP、SP 和 SBL；随后，off-grid SBL 将分数时延与分数多普勒作为超参数迭代更新[13,14]。为兼顾精度和计算量，粗搜索—局部细化的谱匹配框架以低分辨率候选覆盖真实峰，再在局部区域精化[15]；连续域原子范数方法则直接通过二维超分辨模型缓解离散网格误差[16]。这些工作共同推动接收机由“在固定字典中寻找目标”转向“校正字典、分数参数和搜索区域”。")
    add_body(doc, "算法演进随后进一步触及环境先验和网络协同。一方面，环境感知辅助 OTFS 信道估计开始利用被动雷达和 radio map 提取时延、多普勒先验，再执行网格校准和路径细化[10]；另一方面，协同 ISAC 已形成“多站局部参数估计—数据关联—位置速度融合”的处理链[17]，网络化 ISAC 又把多基站波束、UAV 轨迹和基站关联纳入联合优化[18]。这意味着目标状态不再只是最终需要估计的结果，而正在直接参与多个接收机的字典构造、搜索空间选择和网络决策。")
    add_body(doc, "上述发展同时引入了一个此前较少被单独处理的问题：当多个节点依赖同一目标状态构造接收模型时，共享状态一旦偏移，误差将以公共模型失配的形式同时作用于多个接收机。各节点热噪声可以保持独立，但其 delay、Doppler 和 angle 匹配中心会发生方向相关的系统性偏离。因此，增加节点能够平均独立观测波动，却不能保证把错误字典重新移动到真实目标子空间。本文的研究问题由此产生：共享目标状态失配如何限制多节点融合增益，以及能否利用多节点共同观测反向校正这一低维状态，同时避免状态搜索对最终检测统计量产生自拟合偏差。")
    add_heading(doc, "1.2 研究意义", 2)
    add_body(doc, "在理论层面，本文补充多节点空间分集成立所依赖的模型正确性条件。多节点融合能够降低独立测量噪声和节点选择波动，但不必然消除由同一共享状态引起的公共模型误差。通过区分观测不确定性和模型不确定性，可以更严格地解释“增加 UAV 在什么条件下有效、在什么条件下只是重复累积弱匹配证据”。")
    add_body(doc, "在方法层面，本文将研究对象由终端检测概率前移到“状态误差—几何映射—字典子空间偏移—有效信息损失—检测性能变化”的机制链。字典子空间信息保留率用于解释低检测率究竟来自目标能量不足，还是检测器在错误子空间中提取统计量；共享偏移校正则利用多个接收机对同一低维状态的共同约束统一更新字典，而不是分别修补每个节点的局部参数。")
    add_body(doc, "在统计层面，本文明确区分“状态校正确实改善 held-out 检测”和“同一数据既搜索状态又抬高检测统计量”。交叉拟合在本文中只是数据隔离接口，而不是独立的 OTFS 算法创新：一折用于估计共享偏移，另一折用于检测，随后交换角色，检测门限仍由独立 calibration 数据标定。")
    add_body(doc, "在系统层面，共享信息具有双重作用：可信状态可以缩小搜索范围、提高协同效率，错误状态则可能把局部误差转化为网络级公共偏差。因此，合理的系统优化顺序应是先验证接收模型可信度，必要时校正共享状态，再讨论节点数量、功率、波束和队形优化。本文以 6 UAV、3 目标作为一般模型的具体案例，但结论边界限定于短时共享状态失配与检测，不扩展为长时间轨迹跟踪或任意环境先验。")

    add_heading(doc, "2 国内外研究综述", 1)
    add_heading(doc, "2.1 OTFS 感知参数估计：从离散网格到连续域校正", 2)
    add_body(doc, "OTFS 感知参数估计的主线是不断降低接收字典与连续物理参数之间的不一致。早期匹配滤波和近似最大似然方法直接在 DD 网格上估计距离与速度[2]；稀疏恢复方法利用 DD 域少量有效散射路径，采用 OMP、SP 或 SBL 生成目标轮廓[14]。这类方法提高了稀疏场景中的分辨能力，但有限网格仍会造成分数时延和分数多普勒泄漏。")
    add_body(doc, "针对 off-grid 问题，Wei 等把整数网格分量与网格外偏移分离，并通过 SBL 和 EM 联合更新超参数[13]。Xia 等采用粗搜索与三角谱匹配、局部细化与二次谱匹配以及 CFAR 检测的三级流程，在减少全局高分辨率穷举的同时实现检测和参数估计[15]。Liu 等进一步利用二维 atomic norm 在连续 DD 域建立超分辨问题，并通过 SDP 与 ADMM 求解[16]。这些方法分别从概率稀疏建模、粗细搜索和连续域优化缓解网格失配，但通常默认搜索域或先验邻域能够覆盖真实参数。")
    add_heading(doc, "2.2 多目标感知：从独立估计到迭代重构", 2)
    add_body(doc, "多目标场景使误差来源由参数离散化扩展为目标间干扰和迭代误差传播。国内研究采用最大似然参数估计与并行干扰消除，根据上一轮估计重构目标回波、消除已估目标并再次估计，以降低回波叠加造成的误差累积[8]。近期 MIMO-OTFS 方法又根据多目标是否可分选择不同路径：可分目标使用快速二阶反演，不可分目标使用粒子群超分辨搜索处理分数参数[19]。这些工作处理的是目标之间的信号污染或参数不可分辨性；本文关注的是不同节点共同使用错误状态后产生的公共字典偏移，二者不能混为同一类干扰。")
    add_heading(doc, "2.3 网络化协同感知：从空间分集到关联融合与几何优化", 2)
    add_body(doc, "分布式 MIMO 雷达表明，不同收发视角能够利用目标散射差异获得空间分集，但收益依赖同步、几何、融合方式和统计模型[5]。在低空 cooperative ISAC 中，Tang 等进一步构造两阶段流程：各基站先利用空间平滑张量分解和低维 GRQ 提取局部目标参数，再通过最小生成树数据关联、Pareto 位置估计和残差加权融合速度[17]。这类方法把多节点协同由统计量叠加推进到显式的参数关联与状态融合。")
    add_body(doc, "网络化 ISAC 随后把多基站协调波束、UAV 轨迹和基站关联联合起来，并利用 AO、SCA 与 SDR 求解通信速率和感知约束之间的非凸优化[18]。国内研究也开始面向低空无线网络联合优化 UAV 集群部署与波束赋形[12]。这些研究回答了如何获取更多、更高质量的空间观测，但通常把接收模型和共享先验作为可信输入；若字典本身已经偏移，继续增加照射功率或优化几何不一定能够恢复正确的目标子空间。")
    add_heading(doc, "2.4 环境与状态先验辅助感知及其失配", 2)
    add_body(doc, "环境辅助估计使本文问题更加直接。Qing 等利用被动 UAV 雷达估计地面用户位置和速度，通过 radio map 获取传播环境，再从中提取 delay 和 Doppler 先验，提出 sensing-assisted grid calibration 和 path-refinement enhanced channel estimation[10]。该路线证明环境状态可以帮助缩小搜索范围、降低 off-grid 误差，但也意味着环境状态误差会进入网格中心和路径模型。现有研究主要评估辅助先验准确时的增益或对参数扰动的稳健性，较少将“多个接收机共享同一个偏移状态”作为独立误差结构研究。")
    add_heading(doc, "2.5 现有研究缺口与本文定位", 2)
    add_body(doc, "综上，现有文献已经分别解决了固定网格分辨率、off-grid 参数、多目标干扰、多站关联融合以及网络几何和资源优化问题，算法链正在由观测驱动估计走向环境先验辅助和网络级状态共享。然而仍存在三个相连缺口：第一，缺少共享状态偏差如何经多 UAV 几何映射形成相关字典失配的机制量；第二，缺少“增加节点只能降低独立波动、不能自动修复公共模型误差”的适用边界；第三，缺少在不使用真值且不重复使用检测数据的条件下进行共享状态校正的完整检测接口。")
    add_body(doc, "本文据此定位为共享状态失配机制与校正研究，而不是另一种通用 OTFS 参数估计器。问题 1 用子空间信息保留率建立状态误差到有效目标信息的因果链；问题 2 利用跨节点共同状态约束执行 MAP 校正，并通过交叉拟合隔离状态估计与检测。cross fitting 只承担统计隔离功能，核心贡献仍限定为共享失配机制、跨节点状态校正及其检测边界。")

    add_heading(doc, "3 本文的主要解决问题与总体研究框架", 1)
    add_status_note(doc, "统一主线：目标状态 → 节点几何参数 → OTFS 字典 → 有效检测信息 → 共享状态校正 → 数据隔离检测。队形设计和多目标实验仅用于验证该主线的适用边界。")
    add_heading(doc, "3.1 研究对象与核心矛盾", 2)
    add_body(doc, "本文研究短时块静态、多节点协同的无人机 OTFS 感知系统。目标在单个相干处理区间内位置近似不变，但具有非零速度；各 UAV 从不同几何视角观测同一目标，并共享由跟踪器或上一个处理区间提供的目标状态信念。研究对象不是孤立的时延或多普勒栅格点，而是能够同时决定各节点时延、多普勒和角度的连续目标状态。")
    add_body(doc, "系统的核心矛盾在于：多节点协同能够增加观测数量，却也可能重复使用同一份有偏状态先验。当共享信念发生偏移时，各节点并非产生彼此独立的估计误差，而是经各自几何映射形成相关字典失配。因此，论文首先解释公共状态误差为何不能被常规融合自动平均，再研究如何利用跨节点的一致几何约束恢复共享状态，同时避免用同一份数据既估计状态又检验目标。")
    add_table(doc, ["问题", "核心问题", "主要输出", "当前状态"], [
        ("问题 1", "共享偏移为何使多节点融合仍然低检测率", "失配模型、信息保留率、节点数边界", "机制实验已完成"),
        ("问题 2", "如何用共享信息校正状态且避免自拟合", "cross fitted MAP、检测曲线、消融", "探索性证据已完成，正式曲线待跑"),
    ], [2.0, 5.2, 4.2, 3.0])
    add_heading(doc, "3.2 多无人机 OTFS 统一观测模型", 2)
    add_body(doc, "设目标 q 的连续状态由位置和速度组成，各节点的双基地时延、多普勒与角度均由该状态和节点位置共同决定。用状态而不是互不相关的参数栅格作为统一变量，可将节点间的共同信息显式写入模型。")
    add_formula(doc, r"\mathbf x_q=[\mathbf p_q^{T},\mathbf v_q^{T}]^{T}", "1")
    add_formula(doc, r"\boldsymbol\xi_{jq}=g_j(\mathbf x_q)=[\tau_{jq},\nu_{jq},\boldsymbol\theta_{jq}^{T}]^{T}", "2")
    add_formula(doc, r"\mathbf y_j=\mathbf X_j\mathbf h_j+\sum_{q=1}^{Q}\mathbf A_{jq}(\boldsymbol\xi_{jq})\boldsymbol\alpha_{jq}+\mathbf n_j", "3")
    add_body(doc, "式中，Xj hj 表示直达径或结构化干扰，Ajq 是由节点几何参数生成的 OTFS 目标字典，alpha jq 为目标复散射系数，nj 为噪声。当前案例使用 6 架 UAV 和 3 个运动目标，但理论变量保留一般的节点数 M 和目标数 Q；具体数量只承担数值验证，不构成理论前提。")

    add_heading(doc, "3.3 共享状态失配及其几何传播", 2)
    add_body(doc, "所有节点使用同一个目标状态信念构造接收字典。令 delta q 为真值相对于信念的共享状态偏差，并以位置先验标准差 sigma p 归一化位置误差。这样，0、50、150、250 和 400 m 只是 sigma p 等于 150 m 时的案例映射，论文结论以无量纲误差 r 为主。")
    add_formula(doc, r"\widehat{\mathbf x}_q=\mathbf x_q-\boldsymbol\delta_q,\qquad r_q=\|\boldsymbol\delta_{p,q}\|_2/\sigma_p", "4")
    add_body(doc, "在信念状态附近作一阶展开，共享状态偏差在节点 j 上转化为不同的时延、多普勒和角度偏差：")
    add_formula(doc, r"\Delta\boldsymbol\xi_{jq}\approx\mathbf J_{jq}\boldsymbol\delta_q,\qquad \mathbf J_{jq}=\left.\frac{\partial g_j}{\partial\mathbf x}\right|_{\widehat{\mathbf x}_q}", "5")
    add_body(doc, "由此得到本文的机制链：共享状态偏差经过节点几何敏感矩阵形成参数偏移，参数偏移进一步造成 OTFS 字典错位，使真实回波在检测子空间中的信息保留率下降，最终可能降低 GLRT 的非中心性和检测概率。前半段“状态偏差到字典信息损失”已有机制实验支持；信息保留率与 GLRT 检测概率之间的定量映射仍是待补理论与实验，不在现阶段写成既定结论。")

    add_heading(doc, "3.4 多节点信息与可辨识性", 2)
    add_body(doc, "共享误差能否被校正，不由节点数量单独决定，而取决于节点几何提供的独立状态信息。对目标 q，可用下式描述先验与多节点局部曲率共同形成的候选状态信息矩阵：")
    add_formula(doc, r"\mathbf I_q=\mathbf P_{\delta,q}^{-1}+\sum_{j=1}^{M}\mathbf J_{jq}^{T}\mathbf W_{jq}\mathbf J_{jq}", "6")
    add_body(doc, "P delta,q 是共享偏移的先验协方差，W jq 表示节点 j 对几何参数的有效观测权重。若不同节点的 Jjq 提供互补方向，Iq 的最小特征值增大，状态在更多方向上可辨；若节点几何近似共线，即使增加节点，新增信息也可能高度冗余。式（6）在本文中是用于组织实验和解释阵形差异的理论诊断量，尚不能作为已验证的性能定理。")

    add_heading(doc, "3.5 共享状态校正与数据隔离检测", 2)
    add_body(doc, "问题 2 在共享状态空间中执行 MAP 校正，而不是令每个节点独立移动自己的目标字典。对于两折独立或条件独立的数据 YA 和 YB，A 折汇总多节点证据估计共享偏移：")
    add_formula(doc, r"\widehat{\boldsymbol\delta}_{A}=\arg\max_{\boldsymbol\delta\in\Omega}\left[\sum_{j=1}^{M}G_j(\mathbf y_{j,A};\boldsymbol\delta)-\tfrac{1}{2}\boldsymbol\delta^T\mathbf P_{\delta}^{-1}\boldsymbol\delta\right]", "7")
    add_body(doc, "随后只在未参与该次状态估计的 B 折上构造检测统计量，并交换两折角色：")
    add_formula(doc, r"T_{CF}=\sum_{j=1}^{M}T(\mathbf y_{j,B};\widehat{\boldsymbol\delta}_{A})+\sum_{j=1}^{M}T(\mathbf y_{j,A};\widehat{\boldsymbol\delta}_{B})", "8")
    add_body(doc, "交叉拟合的作用是隔离状态估计和目标检测，避免搜索过程把同一份噪声拟合成目标证据。它不是独立创新点，而是使共享状态校正能够接受统一虚警门限检验的统计接口。算法只更新接收机可知的字典、保护基和相关协方差，不修改观测、真实回波、真实信道或噪声。")

    add_heading(doc, "3.6 总体处理链与研究边界", 2)
    add_table(doc, ["阶段", "输入与操作", "输出", "论文作用"], [
        ("先验状态", "读取目标位置、速度信念及协方差", "共享 belief", "给出失配起点"),
        ("几何映射", "由 gj 和 Jjq 映射至时延、多普勒与角度", "节点字典参数", "连接状态域与 DD 域"),
        ("机制诊断", "计算信息保留率与节点子集统计", "共享失配损失", "回答问题 1"),
        ("状态校正", "跨节点 MAP 或低复杂度 WLS", "共享偏移估计", "回答问题 2"),
        ("隔离检测", "用另一折数据构造 GLRT 并融合", "PD、PFA、AUC", "检验真实检测收益"),
        ("边界扩展", "改变 r、几何、目标数和速度误差", "适用区间", "验证泛化而非新增主线"),
    ], [2.3, 6.0, 3.1, 3.2])
    add_body(doc, "据此，全文只保留两项核心研究内容：其一，建立共享状态误差经多节点几何传播为相关字典失配的机制模型；其二，建立具有数据隔离约束的共享状态校正方法。可辨识性分析、20 m 节点间距、队形对照和多目标实验均服务于这两项内容，不再扩展为彼此平行的新方向。")

    add_heading(doc, "4 问题 1 共享目标状态失配的机理与边界", 1)
    add_heading(doc, "4.1 字典失配与信息保留率", 2)
    add_body(doc, "在第 3 节统一模型基础上，问题 1 固定 UAV 数量、目标数量、阵形、发射功率、直达径消除和融合规则。节点 j 使用信念状态构造目标字典 Ab,j，而回波由真值字典 At,j 产生。若偏差是公共的，则所有节点的字典中心同时偏离真实时延、多普勒和方位参数。独立噪声可通过融合减小，字典中心的系统偏差不会随节点数增加而消失。")
    add_formula(doc, r"\eta_j(\boldsymbol\delta)=\frac{\|\mathbf Q_{b,j}^{H}\mathbf A_{t,j}\|_F^2}{\|\mathbf A_{t,j}\|_F^2},\quad \mathbf Q_{b,j}=\operatorname{orth}(\mathbf A_{b,j})", "9")
    add_formula(doc, r"\eta_{\mathcal S}=\frac{\sum_{j\in\mathcal S}\eta_j\|\mathbf A_{t,j}\|_F^2}{\sum_{j\in\mathcal S}\|\mathbf A_{t,j}\|_F^2}", "10")
    add_body(doc, "eta j 衡量真值目标字典中有多少能量仍落在信念字典子空间内。eta 等于 1 表示字典完全覆盖真实回波，接近 0 表示检测器在错误子空间中提取统计量。对节点集合 S 的信息保留率是按真实回波能量加权的平均值，因此增加节点可减小节点选择方差，却不会系统性改变公共失配的均值。")

    add_heading(doc, "4.2 优化函数及分析方法", 2)
    add_body(doc, "问题 1 不引入新的迭代优化器，而是采用可重复的机制量。对每个场景和误差半径，分别构造真值目标字典与失配字典，计算 6 个接收机的信息保留率；随后枚举不同节点数的全部接收机子集，计算子集均值和子集间标准差。该分析避免检测门限、噪声实现和状态搜索共同影响结果。")
    add_body(doc, "实验配置为 6 架 UAV、3 个目标、12 个独立场景、最短三维节点间距 20 m、位置先验标准差 150 m。误差半径取 0、50、150、250 和 400 m，对应 r 等于 0、0.33、1、1.67 和 2.67。400 m 仅作为失配压力测试，不代表常规跟踪误差。")

    add_heading(doc, "4.3 实验结果及分析", 2)
    mech_rows = []
    for radius in (0.0, 50.0, 150.0, 250.0, 400.0):
        mean, low, high = summary_value(retention, radius)
        mech_rows.append((f"{radius:.0f}", f"{radius / 150:.2f}",
                          f"{mean:.3f}", f"[{low:.3f}, {high:.3f}]"))
    add_table(doc, ["位置误差 m", "归一化误差 r", "六节点信息保留率", "场景分位区间 2.5% 到 97.5%"], mech_rows, [3.0, 3.0, 4.0, 4.4])
    add_body(doc, "无误差时六节点信息保留率为 1.000。误差为 50 m 时均值降至 0.716；误差达到一个先验标准差 150 m 时仅保留 0.249；250 m 和 400 m 时进一步降至 0.097 和 0.039。场景间区间较宽，说明几何会调节失配敏感性，但总体下降趋势清晰。")
    one_150 = summary_value(retention, 150.0, 1)[0]
    six_150 = summary_value(retention, 150.0, 6)[0]
    one_400 = summary_value(retention, 400.0, 1)[0]
    six_400 = summary_value(retention, 400.0, 6)[0]
    add_body(doc, f"节点数没有恢复平均信息。在 150 m 误差下，1 节点与 6 节点的平均保留率分别为 {one_150:.3f} 和 {six_150:.3f}；在 400 m 误差下分别为 {one_400:.3f} 和 {six_400:.3f}。六节点的子集选择方差归零是因为全部节点均被选中，并不表示公共偏差得到修正。该结果支持“融合减少独立波动，但不能消除共享字典偏差”的机制判断。")
    add_body(doc, "代表性旧场景中，瓶颈目标的位置偏差达到 531 m，六个单节点检测概率均低于 0.18，融合后约为 0.263；将 belief 临时替换为真值后，同一系统的融合检测概率接近 1。该案例与信息保留率实验方向一致，但它属于单场景反事实审计，不作为总体性能估计。")

    add_heading(doc, "4.4 消融实验及分析", 2)
    add_table(doc, ["消融因素", "设计", "判别逻辑", "状态"], [
        ("节点数", "M 取 1 到 6，枚举全部子集", "均值不升且方差下降，支持公共偏差解释", "已完成"),
        ("误差尺度", "r 取 0 到 2.67", "检验信息保留率是否随失配增大而下降", "已完成"),
        ("几何", "固定随机几何并满足 20 m 间距", "排除非法近距离节点", "已完成"),
        ("独立误差对照", "各节点注入独立方向误差", "与共享偏差比较融合可平均性", "待补"),
        ("速度误差", "固定位置误差，改变速度先验", "检验 OTFS 多普勒维失配", "待补"),
    ], [2.8, 4.2, 5.1, 2.3])

    add_heading(doc, "4.5 实验总结", 2)
    add_body(doc, "问题 1 已形成可写入论文的初步机制结论：共享目标状态偏差会使所有节点的 OTFS 目标字典同时偏离真值；增加节点主要降低节点选择造成的波动，不恢复平均字典信息。该结论目前由 12 场景子空间实验和一个代表性 oracle 反事实案例支持。其边界是信息保留率并非检测概率，仍需在统一门限与独立测试集下建立 eta 与 AUC、PD 的对应关系。")

    add_heading(doc, "5 问题 2 共享状态的交叉拟合校正", 1)
    add_heading(doc, "5.1 两折数据与共享状态校正", 2)
    add_body(doc, "问题 2 沿用第 3 节式（7）和式（8）的共享 MAP 与交叉拟合检测模型。假设相邻两个短时相干处理区间内目标位置偏移近似不变，记两折数据为 YA 和 YB，每一折均包含 6 个接收机的观测。算法使用 YA 估计共享二维位置偏移并只在 YB 上检测，随后交换角色。该假设要求折间状态漂移小于目标字典的有效相关宽度，不能扩展为长时间静态假设。")
    add_body(doc, "式（7）中的 Gj 是接收机 j 对候选偏移的状态拟合目标，P delta 是共享偏移先验协方差，Omega 是搜索域；式（8）中的 T 是 held out 检测统计量。实现通过字段白名单保证只更新接收机可知的字典、保护基和相关协方差，从而防止真值泄漏。")

    add_heading(doc, "5.2 优化函数及优化算法", 2)
    add_body(doc, "当前基线求解器采用粗到细 MAP 搜索。已有峰宽实验表明，目标函数的半高宽约为 25 至 75 m，因此粗网格间距采用 75 m，而不是早期方案中的 300 m。候选偏移先由多接收机目标函数求和评分，再围绕最优候选进行局部细化。正式论文保留这一冻结基线，不再把 Top K、NMS、边界扩展等机制同时加入主方法。")
    add_bullets(doc, [
        "名义方法：不估计偏移，直接使用原始 belief。",
        "self fit：用同一折估计偏移并检测同一折，用于暴露重复使用数据的偏差。",
        "local cross fit：每个接收机独立估计偏移后交叉检测，用于检验共享建模的贡献。",
        "shared cross fit：全部接收机共同估计一个偏移并交叉检测，是本文拟议方法。",
        "oracle：使用真值状态，只作为不可部署上界。",
    ])
    add_body(doc, "低复杂度 Jacobian WLS MAP 已实现并完成小样本筛查。其单场景求解时间约为 0.35 s，显著低于网格搜索约 16 s；但在 150 m 误差的 3 个试验场景中，WLS 残余状态误差约为 126 至 152 m，而网格 MAP 为 2 至 17 m。因此 WLS 目前只能作为速度基线，不能替代主方法。")

    add_heading(doc, "5.3 实验结果及分析", 2)
    nominal = eval_by_name["nominal"]
    local = eval_by_name["local_crossfit"]
    shared = eval_by_name["shared_crossfit"]
    oracle = eval_by_name["oracle"]
    add_status_note(doc, "以下 Gate 1 是 150 m 单档探索性证据。其场景数有限，且生成于 20 m 最短间距成为正式约束之前，不能作为最终确认性结果。")
    add_table(doc, ["方法", "AUC", "PD", "经验 PFA", "状态误差"], [
        ("nominal", f"{nominal['test_auc']:.4f}", f"{nominal['test_pd']:.2f}", f"{nominal['empirical_pfa']:.2f}", "150.0 m 基线"),
        ("local cross fit", f"{local['test_auc']:.4f}", f"{local['test_pd']:.2f}", f"{local['empirical_pfa']:.2f}", "各节点独立"),
        ("shared cross fit", f"{shared['test_auc']:.4f}", f"{shared['test_pd']:.2f}", f"{shared['empirical_pfa']:.2f}", f"中位 {shared['state_median_m']:.1f} m"),
        ("oracle", f"{oracle['test_auc']:.4f}", f"{oracle['test_pd']:.2f}", f"{oracle['empirical_pfa']:.2f}", "0 m"),
    ], [3.4, 2.4, 2.4, 2.6, 3.6])
    add_body(doc, f"在该探索性数据中，shared cross fit 相对 nominal 的 AUC 增量为 {shared['delta_auc_vs_nominal']:.4f}，PD 增量为 {shared['delta_pd_vs_nominal']:.2f}。AUC 增量的配对 95% 区间为 [{shared['delta_auc_ci95'][0]:.4f}, {shared['delta_auc_ci95'][1]:.4f}]，跨越 0；PD 增量区间下界为 {shared['delta_pd_ci95'][0]:.2f}。因此可以写成“显示有利趋势”，不能写成“已经显著优于”。")
    add_body(doc, "shared cross fit 的 AUC 高于 local cross fit，且经验 PFA 为 0.05；local cross fit 的经验 PFA 为 0.20。这一结果表明共享状态假设可能提供独立贡献。不过每种方法只有 20 个测试 H0 样本，经验 PFA 只能以 0.05 为步长变化，置信区间较宽。正式实验必须扩大标定和测试规模，并在新几何约束下重跑。")
    if stageb is not None:
        audit = stageb["audit"]
        diag = {d["fold"]: d for d in stageb.get("diagnosis", [])}
        add_status_note(doc, "阶段 B 的 r=1 初始冒烟未通过；随后只增加候选覆盖的单场景反事实已恢复两折状态，尚待 r=0 门禁标定后扩展。")
        add_table(doc, ["检查项", "结果", "裁决"], [
            ("墙钟时间", f"{audit['runtime_s']:.1f} s", "记录成本"),
            ("最小节点间距", f"{audit['minimum_uav_separation_m']:.1f} m", "通过 20 m 约束"),
            ("shared 平均残余误差", f"{audit['shared_state_error_mean_m']:.1f} m", "高于 150 m 基线，未通过"),
            ("WLS 平均残余误差", f"{audit['wls_state_error_mean_m']:.1f} m", "该场景向真值改善"),
            ("A 折诊断", (f"J真值={diag['A']['objective_truth']:.2f}，"
                           f"J选中={diag['A']['objective_selected']:.2f}"
                           if "A" in diag else "待诊断"), "搜索漏掉更优真峰"),
            ("B 折诊断", (f"J真值={diag['B']['objective_truth']:.2f}，"
                           f"J选中={diag['B']['objective_selected']:.2f}"
                           if "B" in diag else "待诊断"), "目标面偏好假峰"),
            ("37.5 m 候选覆盖 A 折", (f"残余误差 {diag['A']['dense_best_error_m']:.1f} m"
                                      if "A" in diag and "dense_best_error_m" in diag["A"]
                                      else "待诊断"), "候选覆盖修复成立"),
            ("37.5 m 候选覆盖 B 折", (f"残余误差 {diag['B']['dense_best_error_m']:.1f} m"
                                      if "B" in diag and "dense_best_error_m" in diag["B"]
                                      else "待诊断"), "更新有效，无需回退"),
        ], [4.0, 5.4, 5.0])
        add_body(doc, "初始反例说明失败不能统一归因于搜索边界。A 折真值附近的 75 m 粗网格候选仅排第 11；将候选间距减半至 37.5 m 后，A、B 两折残余误差分别降至 16.7 m 和 39.7 m。B 折虽然没有选择精确真值点，但其更新具有正目标增益、4 个接收机支持、不贴边，且与 A 折只相差 37.5 m，因此不应事后按真值强制回退。正式门禁阈值必须由 r=0 calibration 场景标定。")

    add_heading(doc, "5.4 消融实验及分析", 2)
    add_table(doc, ["消融", "回答的问题", "主指标", "判定"], [
        ("self fit 对 cross fit", "数据隔离是否控制 H0 统计量抬升", "PFA、H0 分布、AUC", "各臂独立标定门限"),
        ("local 对 shared", "共享状态建模是否具有独立贡献", "配对 AUC、PD、状态误差", "shared 优于 local"),
        ("nominal 对 oracle", "方法可闭合多少失配上限", "oracle gap closure", "只解释可恢复空间"),
        ("WLS 对 grid MAP", "复杂度与鲁棒性的交换", "运行时间、RMSE、失败率", "不只比较平均速度"),
        ("搜索步长", "求解失败还是目标面不可辨识", "真值处目标值、峰宽、恢复误差", "步长不大于峰宽"),
        ("误差档位", "校正在什么失配区间有效", "r 分层 AUC、PD、PFA", "主轴采用 r"),
    ], [3.0, 5.0, 3.3, 3.2])
    add_body(doc, "此前 gated Top K 消融的平均状态误差改善主要由少数场景贡献，去掉两个最大改善场景后均值改善消失。因此该机制不进入本文主方法。边界假峰目前只是候选解释，不是已证实病因；只有当失败场景在预先定义的边界诊断量上稳定聚集，并能被单因素修复时，才能形成病因结论。")

    add_heading(doc, "5.5 实验总结", 2)
    add_body(doc, "问题 2 已具备完整的算法实现、信息隔离测试和一档探索性检测结果。现有证据支持 shared cross fit 具有恢复状态和改善检测的潜力，也显示全局网格搜索存在较高计算成本。当前不足不是继续叠加求解器模块，而是缺少 20 m 节点间距、归一化误差分层和足够标定样本下的确认性曲线。")

    add_heading(doc, "6 文章总体总结与接下来的优化操作", 1)
    add_heading(doc, "6.1 总体总结", 2)
    add_body(doc, "本文已从分散的队形、接收机和状态搜索尝试中收缩出一条可检验主线。问题 1 解释公共目标状态误差为何限制多节点融合，问题 2 用数据隔离的共享 MAP 更新恢复目标字典。OTFS 的作用落实在运动目标的时延多普勒联合表征，6 UAV 与 3 目标是该一般模型的一个具体案例。误差不再只用米表示，而以 r 等于偏差除以先验标准差表征，再映射到工程尺度。")
    add_body(doc, "当前最可靠的新证据是 12 场景信息保留率曲线。150 m 误差即 r 等于 1 时，六节点平均只保留约四分之一的真值字典信息；400 m 压力误差下约为 4%。探索性 Gate 1 显示共享交叉拟合可能闭合部分 oracle 缺口，但尚未达到确认性证据强度。论文结论必须保持这一区分。")

    add_heading(doc, "6.2 分阶段实验计划", 2)
    add_table(doc, ["阶段", "实验", "最小产出", "继续条件"], [
        ("A 已完成", "共享失配信息保留率", "12 场景，r 分层，M 等于 1 至 6", "趋势单调且 M 不恢复均值"),
        ("B 单场景已修复", "正式单目标检测曲线", "r 等于 0、0.33、1、1.67；20 m 间距", "先用 r=0 标定门禁再扩样"),
        ("C", "数据隔离与共享性消融", "self、local、shared、oracle", "PFA 受控且 shared 优于 local"),
        ("D", "三目标泛化", "逐目标 PD 与 min q PD", "优势不是单一目标特例"),
        ("E 可选", "belief aware 队形", "随机、单视角、双视角对照", "仅当 B 到 D 通过"),
    ], [2.1, 4.5, 5.0, 3.0])
    add_body(doc, "阶段 B 的正式误差档采用 r 等于 0、0.33、1 和 1.67。若 sigma p 等于 150 m，对应 0、50、150 和 250 m。400 m 即 r 等于 2.67 只作为压力测试，不能与常规工作区混合汇总。每个误差档分别用 calibration 集的 H0 样本标定门限，AUC 与 PD 只在 test 集计算。主终点为配对 AUC 差；PD、PFA、状态中位误差、RMSE 和运行时间为次终点。")
    add_heading(doc, "6.3 统计与停止规则", 2)
    add_bullets(doc, [
        "正式实验前先用 8 至 12 个场景估计配对差标准差，再进行样本量计算。",
        "任何新机制若使所需场景数超过 100，先修复方差来源，不直接扩大实验。",
        "所有方法使用相同场景、随机种子和几何，报告配对差与 bootstrap 区间。",
        "PFA 门限只使用 calibration 的 H0；test 数据不得参与门限或求解器选择。",
        "若 shared 未稳定优于 local，则删去共享建模创新点，只保留状态失配诊断。",
        "若问题 2 未通过，不启动队形优化正式实验。",
    ])
    add_heading(doc, "6.4 近期实现操作", 2)
    add_table(doc, ["优先级", "操作", "实现位置", "预计成本"], [
        ("P0", "冻结 coarse to fine 共享 cross fit 为主方法", "target_state fit 与 crossfit", "无新增运行成本"),
        ("P0", "建立新的正式分层驱动并锁定 20 m 约束", "tools 与 direction4 data", "开发与冒烟约半天"),
        ("P0", "先跑每档 8 至 12 场景估计方差", "r 等于 0、0.33、1、1.67", "约 2 至 4 小时"),
        ("P1", "按功效分析扩展确认性样本", "仅扩展通过冒烟的档位", "由实测方差决定"),
        ("P1", "补独立节点误差与速度误差机制消融", "问题 1 扩展", "分钟到小时级"),
        ("P2", "三目标 min q PD 泛化", "问题 2 扩展", "确认性曲线通过后"),
    ], [1.8, 5.3, 4.3, 3.2])

    add_heading(doc, "参考文献", 1)
    references = [
        "[1] Liu F, Cui Y, Masouros C, et al. Integrated Sensing and Communications: Toward Dual-Functional Wireless Networks for 6G and Beyond. IEEE Journal on Selected Areas in Communications, 2022, 40(6): 1728-1767. DOI: 10.1109/JSAC.2022.3156632.",
        "[2] Gaudio L, Kobayashi M, Caire G, Colavolpe G. On the Effectiveness of OTFS for Joint Radar Parameter Estimation and Communication. IEEE Transactions on Wireless Communications, 2020, 19(9): 5951-5965. DOI: 10.1109/TWC.2020.2998583.",
        "[3] Shtaiwi E, Abdelhadi A, Li H, et al. Orthogonal Time Frequency Space for Integrated Sensing and Communication: A Survey. arXiv:2402.09637, 2024.",
        "[4] Meng K, Wu Q, Xu J, et al. UAV-Enabled Integrated Sensing and Communication: Opportunities and Challenges. arXiv:2206.03408, 2022.",
        "[5] Haimovich A M, Blum R S, Cimini L J. MIMO Radar with Widely Separated Antennas. IEEE Signal Processing Magazine, 2008, 25(1): 116-129. DOI: 10.1109/MSP.2008.4408448.",
        "[6] Chernozhukov V, Chetverikov D, Demirer M, et al. Double Debiased Machine Learning for Treatment and Structural Parameters. The Econometrics Journal, 2018, 21(1): C1-C68. DOI: 10.1111/ectj.12097.",
        "[7] Nordio A, Chiasserini C F, Viterbo E. Joint Communication and Sensing in OTFS-Based UAV Networks. arXiv:2311.17742, 2023; IEEE Xplore document 10891196.",
        "[8] 陈佳彬, 王朝炜, 庞明亮, 等. 基于 OTFS 的通感一体化主动信道感知与低空多目标探测. 物联网学报, 2024, 8(3).",
        "[9] 赵千禧, 刘嘉宁, 王帝文, 田峰. 基于 OTFS 的通感一体化感知目标参数估计新算法. 数据采集与处理, 2025. DOI: 10.16337/j.1004-9037.2025.06.003.",
        "[10] Qing C, Liu Z, Hu W, et al. Environmental Sensing-Assisted Off-Grid Channel Estimation in UAV-Enabled OTFS-ISAC Systems. Chinese Journal of Aeronautics, 2026, 39(7). DOI: 10.1016/j.cja.2025.103923.",
        "[11] 杨小龙, 于子亮, 张毅, 等. 基于 DD 域 LFM 脉冲信号嵌入的 OTFS 通感一体化波形设计. 电子学报, 2026: 1-9. DOI: 10.12263/DZXB.20260323.",
        "[12] 毛炜昊, 陆杨, 等. 面向低空无线网络的网络化 ISAC 无人机集群部署和波束赋形协同优化. 电子学报, 2026, 54(6): 204-216. DOI: 10.12263/DZXB.20260121.",
        "[13] Wei Z, Li W, Yuan J, Ng D W K. Off-Grid Channel Estimation With Sparse Bayesian Learning for OTFS Systems. IEEE Transactions on Wireless Communications, 2022, 21(9): 7407-7426. DOI: 10.1109/TWC.2022.3158616.",
        "[14] Zacharia O, Devi M V. Target Parameter Estimation for OTFS Radar Using Sparse Signal Processing. Physical Communication, 2023, 58: 102040. DOI: 10.1016/j.phycom.2023.102040.",
        "[15] Xia X, et al. Achieving Better Accuracy With Less Computations: A Delay-Doppler Spectrum Matching Assisted Active Sensing Framework for OTFS Based ISAC Systems. IEEE Transactions on Wireless Communications, 2024, 23(6): 6204-6220. DOI: 10.1109/TWC.2023.3330845.",
        "[16] Liu S, Zhang H, Li L, et al. Super-Resolution Delay-Doppler Estimation for OTFS-Based Automotive Radar. Signal Processing, 2024, 224: 109596. DOI: 10.1016/j.sigpro.2024.109596.",
        "[17] Tang J, Yu Y, Pan C, et al. Cooperative ISAC-Empowered Low-Altitude Economy. IEEE Transactions on Wireless Communications, 2025: 3837-3853. DOI: 10.1109/TWC.2025.3542399.",
        "[18] Cheng G, Song X, Lyu Z, Xu J. Networked ISAC for Low-Altitude Economy: Coordinated Transmit Beamforming and UAV Trajectory Design. IEEE Transactions on Communications, 2025, 73(8): 5832-5847. DOI: 10.1109/TCOMM.2025.3541027.",
        "[19] Integrated Sensing and Communication With MIMO-OTFS: Energy-Conscious Channel Reconstruction via Efficient Target Parameter Estimation. IEEE Transactions on Green Communications and Networking, 2026, 10: 1552-1564. DOI: 10.1109/TGCN.2025.3641228.",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(ref))

    doc.add_page_break()
    add_heading(doc, "附录 A 当前证据与论文表述边界", 1)
    add_table(doc, ["证据", "可写表述", "不可写表述"], [
        ("12 场景信息保留率", "共享失配随 r 增大显著削弱字典重合", "已经证明 PD 必然按相同比例下降"),
        ("Gate 1 150 m", "shared cross fit 显示 AUC 和 PD 有利趋势", "方法已显著优于全部基线"),
        ("单场景 oracle", "该异常场景的低 PD 主要由 belief 失配解释", "所有低 PD 都由同一原因产生"),
        ("WLS 小样本", "WLS 更快但该试验中状态恢复不足", "WLS 普遍无效"),
        ("Top K 配对消融", "改善对少数场景敏感", "边界假峰已被证实为唯一病因"),
    ], [3.2, 5.7, 5.5])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("多无人机 OTFS 通感一体化共享状态失配研究阶段报告"), 9, color=(90, 90, 90))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
